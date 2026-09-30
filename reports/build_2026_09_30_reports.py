"""Render the reviewed LTGD Markdown reports and their deterministic diagrams."""

from __future__ import annotations

import json
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
FIGURES = ROOT / "figures"
FIGURES.mkdir(exist_ok=True)
FONT = Path("C:/Windows/Fonts/arial.ttf")
BOLD = Path("C:/Windows/Fonts/arialbd.ttf")
NAVY = (31, 55, 78)
BLUE = (46, 95, 133)
PALE = (235, 243, 248)
LIGHT = (247, 249, 251)
GRAY = (83, 93, 104)
GREEN = (38, 115, 89)
RED = (156, 67, 55)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(BOLD if bold else FONT), size)


def canvas(width: int, height: int):
    im = Image.new("RGB", (width, height), "white")
    return im, ImageDraw.Draw(im)


def box(draw, xy, title, lines=(), fill=PALE, title_size=38, body_size=29):
    draw.rounded_rectangle(xy, radius=23, fill=fill, outline=NAVY, width=4)
    x0, y0, x1, y1 = xy
    y = y0 + 23
    for line in title.split("\n"):
        draw.text((x0 + 24, y), line, fill=NAVY, font=font(title_size, True))
        y += title_size + 4
    y += 7
    for line in lines:
        draw.text((x0 + 24, y), line, fill=GRAY, font=font(body_size))
        y += body_size + 8


def arrow(draw, points, color=NAVY, width=6, label=None, label_xy=None, label_size=26):
    draw.line(points, fill=color, width=width, joint="curve")
    x0, y0 = points[-2]
    x1, y1 = points[-1]
    size = 18
    if x1 > x0:
        tri = [(x1, y1), (x1 - size, y1 - 10), (x1 - size, y1 + 10)]
    elif x1 < x0:
        tri = [(x1, y1), (x1 + size, y1 - 10), (x1 + size, y1 + 10)]
    elif y1 > y0:
        tri = [(x1, y1), (x1 - 10, y1 - size), (x1 + 10, y1 - size)]
    else:
        tri = [(x1, y1), (x1 - 10, y1 + size), (x1 + 10, y1 + size)]
    draw.polygon(tri, fill=color)
    if label and label_xy:
        draw.text(label_xy, label, fill=color, font=font(label_size, True))


def architecture():
    im, d = canvas(1800, 620)
    d.text((60, 30), "Runtime and evidence boundaries", fill=NAVY, font=font(44, True))
    box(d, (60, 140, 360, 350), "Pi conversation", ["Natural-language", "game request"], title_size=31)
    box(d, (450, 140, 780, 350), "Generator", ["Pi file tools", "write project"])
    box(d, (870, 140, 1190, 350), "Godot project", ["Selected output", "directory"])
    box(d, (1280, 140, 1740, 350), "LTGD Executor", ["End-of-turn hook", "binds declared path"])
    arrow(d, [(360, 245), (450, 245)])
    arrow(d, [(780, 245), (870, 245)])
    arrow(d, [(1190, 245), (1280, 245)])
    box(d, (260, 445, 760, 575), "Planner model", ["Only after confirmed failure"], fill=LIGHT, title_size=32, body_size=25)
    box(d, (1000, 445, 1530, 575), "Executor checks", ["Godot local; review uses model"], fill=LIGHT, title_size=32, body_size=25)
    arrow(d, [(1470, 350), (1470, 420), (1260, 420), (1260, 445)], color=BLUE)
    arrow(d, [(1000, 510), (760, 510)], color=RED, label="failure", label_xy=(825, 470), label_size=24)
    arrow(d, [(500, 445), (500, 390), (610, 390), (610, 350)], color=GREEN, label="repair steps", label_xy=(320, 390), label_size=22)
    im.save(FIGURES / "ltgd_architecture_2026-09-30.png")


def control_flow():
    im, d = canvas(1800, 860)
    d.text((60, 28), "Failure-triggered calls and token costs", fill=NAVY, font=font(44, True))
    box(d, (60, 125, 385, 325), "Generator", ["Write game", "handoff path"])
    box(d, (485, 125, 810, 325), "Executor", ["Godot checks", "0 LLM tokens"])
    box(d, (910, 125, 1235, 325), "Executor review", ["Extra model call", "reads project files"], title_size=31)
    box(d, (1360, 125, 1735, 325), "Done", ["No plan on PASS"], fill=(231, 246, 237))
    arrow(d, [(385, 225), (485, 225)])
    arrow(d, [(810, 225), (910, 225)], color=GREEN, label="PASS", label_xy=(828, 175))
    arrow(d, [(1235, 225), (1360, 225)], color=GREEN, label="MET", label_xy=(1265, 175), label_size=23)
    box(d, (700, 535, 1070, 735), "Planner", ["Model call on failure", "scoped repair evidence"], fill=LIGHT)
    box(d, (1360, 535, 1735, 735), "Stopped", ["Cannot repair or", "invalid handoff"], fill=(250, 242, 240))
    arrow(d, [(645, 325), (645, 635), (700, 635)], color=RED, label="GODOT FAIL", label_xy=(460, 435), label_size=23)
    arrow(d, [(1070, 325), (1070, 470), (885, 470), (885, 535)], color=RED, label="MISSING", label_xy=(1092, 405), label_size=23)
    arrow(d, [(700, 675), (220, 675), (220, 325)], color=BLUE, label="REVISE", label_xy=(365, 630), label_size=23)
    arrow(d, [(1070, 635), (1360, 635)], color=RED, label="CANNOT RESOLVE", label_xy=(1090, 590), label_size=23)
    d.text((65, 790), "Review adds token cost on Godot PASS; repair adds Planner and Generator calls.", fill=GRAY, font=font(25))
    im.save(FIGURES / "ltgd_control_flow_2026-09-30.png")


def legacy_flow():
    im, d = canvas(1800, 550)
    d.text((60, 28), "Earlier report Figure 2  |  relationship redrawn", fill=NAVY, font=font(42, True))
    top = [(60, 125, 380, 290), (490, 125, 810, 290), (920, 125, 1240, 290), (1350, 125, 1730, 290)]
    box(d, top[0], "Direct attempt", ["Pi edits game"], title_size=32, body_size=25)
    box(d, top[1], "Verify", ["Import and boot"], title_size=32, body_size=25)
    box(d, top[2], "Verified", ["Technical pass"], title_size=32, body_size=25)
    box(d, top[3], "Finish gate", ["godot_finish"], title_size=32, body_size=25)
    for a, b in zip(top, top[1:]):
        arrow(d, [(a[2], 205), (b[0], 205)])
    box(d, (305, 380, 680, 515), "Local repair", ["First failure"], fill=LIGHT, title_size=29, body_size=24)
    box(d, (750, 380, 1120, 515), "Verify again", ["Second check"], fill=LIGHT, title_size=29, body_size=24)
    box(d, (1190, 380, 1660, 515), "Plan and act", ["Only after second failure"], fill=LIGHT, title_size=29, body_size=24)
    arrow(d, [(650, 290), (650, 350), (492, 350), (492, 380)], color=RED, label="FAIL", label_xy=(670, 315), label_size=22)
    arrow(d, [(680, 445), (750, 445)])
    arrow(d, [(1120, 445), (1190, 445)], color=RED, label="FAIL", label_xy=(1124, 400), label_size=22)
    im.save(FIGURES / "ltgd_legacy_flow_2026-09-30.png")


def total_tokens_chart():
    usage_root = ROOT.parent / "output" / "token_usage"
    games = [
        ("Horror Signal Lost", "HSL"),
        ("Puzzle Magnet Lab", "MAG"),
        ("Ivory Beats", "PIB"),
    ]
    im, d = canvas(1800, 850)
    d.text((60, 30), "Total tokens by game", fill=NAVY, font=font(44, True))
    d.rectangle((1280, 50, 1314, 84), fill=NAVY)
    d.text((1324, 49), "Baseline", fill=GRAY, font=font(26))
    d.rectangle((1508, 50, 1542, 84), fill=GREEN)
    d.text((1552, 49), "LTGD", fill=GRAY, font=font(26))

    left, right, top, bottom = 185, 1730, 155, 635
    for tick in range(0, 41, 10):
        y = bottom - (tick / 40) * (bottom - top)
        d.line([(left, y), (right, y)], fill=(219, 226, 231), width=2)
        d.text((107 if tick == 0 else 92, y - 16), f"{tick}", fill=GRAY, font=font(24))
    d.text((59, 113), "Million tokens", fill=GRAY, font=font(24))

    centers = [420, 960, 1500]
    bar_width = 150
    for center, (name, code) in zip(centers, games):
        values = []
        for run in ("direct", "LTGD"):
            path = usage_root / f"{code}_{run}" / "usage.json"
            values.append(json.loads(path.read_text(encoding="utf-8"))["totals"]["total"])
        for value, x, color in zip(values, (center - 170, center + 20), (NAVY, GREEN)):
            y = round(bottom - value / 40_000_000 * (bottom - top))
            d.rectangle((x, y, x + bar_width, bottom), fill=color)
            label = f"{value / 1_000_000:.2f}M"
            bbox = d.textbbox((0, 0), label, font=font(25, True))
            d.text((x + (bar_width - (bbox[2] - bbox[0])) / 2, y - 37), label, fill=color, font=font(25, True))
        label_box = d.textbbox((0, 0), name, font=font(26, True))
        d.text((center - (label_box[2] - label_box[0]) / 2, 654), name, fill=NAVY, font=font(26, True))
        saving = (values[0] - values[1]) / values[0] * 100
        saving_label = f"{saving:.1f}% fewer"
        saving_box = d.textbbox((0, 0), saving_label, font=font(24))
        d.text((center - (saving_box[2] - saving_box[0]) / 2, 700), saving_label, fill=GREEN, font=font(24))

    d.text((185, 789), "One archived run per condition; totals include cached reads.", fill=GRAY, font=font(22))
    im.save(FIGURES / "ltgd_total_tokens_2026-09-30.png")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("LTGD  |  ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def clean_inline(text: str) -> str:
    return text.replace("**", "").replace("`", "").replace("<br>", "\n")


def add_table(doc, rows, chinese):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    n = len(rows[0])
    widths = {3: [1.32, 2.2, 3.25], 4: [1.27, 1.77, 1.83, 1.9], 5: [1.77, 1.05, 1.4, 1.4, 1.15], 6: [1.65, 0.70, 1.22, 1.18, 0.92, 1.10], 7: [1.50, 0.69, 0.91, 0.91, 0.78, 1.04, 0.94], 8: [1.43, 0.55, 0.46, 0.76, 1.04, 0.69, 1.02, 0.82]}.get(n, [6.77 / n] * n)
    if n == 4:
        widths = [0.93, 1.88, 2.19, 1.77] if chinese else [1.32, 2.15, 2.15, 1.15]
    for i, value in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.text = clean_inline(value)
        set_cell_shading(cell, "EAF1F6")
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            cells[i].text = clean_inline(value)
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                if ri == 0:
                    p.paragraph_format.keep_with_next = True
                if n in (6, 7, 8) and ri > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER if ci == 1 else WD_ALIGN_PARAGRAPH.RIGHT
                if n == 4 and not chinese and ri > 0 and ci == 3:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Microsoft YaHei" if chinese else "Arial"
                    run.font.size = Pt(8.4 if chinese else 8.2 if n == 8 else 8.7)
                    run.font.bold = ri == 0 or (n == 4 and not chinese and ri > 0 and ci == 3)
                    run.font.color.rgb = RGBColor(*NAVY) if ri == 0 else RGBColor(*GREEN) if n == 4 and not chinese and ci == 3 else RGBColor(35, 42, 48)
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        if ri == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build(markdown_name: str, docx_name: str, chinese: bool):
    lines = (ROOT / markdown_name).read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.73)
    sec.left_margin = sec.right_margin = Inches(0.82)
    add_page_number(sec)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei" if chinese else "Arial"
    normal.font.size = Pt(10.0 if chinese else 10.5)
    normal.font.color.rgb = RGBColor(35, 42, 48)
    normal.paragraph_format.space_after = Pt(5 if chinese else 7)
    normal.paragraph_format.line_spacing = 1.12 if chinese else 1.17
    title = styles["Title"]
    title.font.name = "Microsoft YaHei" if chinese else "Arial"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(7)
    title_ppr = title._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    for name, size in [("Heading 1", 14), ("Heading 2", 11.5)]:
        st = styles[name]
        st.font.name = "Microsoft YaHei" if chinese else "Arial"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0) if chinese else RGBColor(*NAVY)
        st.paragraph_format.space_before = Pt(10 if chinese else 13)
        st.paragraph_format.space_after = Pt(5 if chinese else 6)
        st.paragraph_format.keep_with_next = True
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            doc.add_paragraph(clean_inline(line[2:]), style="Title")
        elif line.startswith("## "):
            doc.add_paragraph(clean_inline(line[3:]), style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(clean_inline(line[4:]), style="Heading 2")
        elif line.startswith("!["):
            match = re.match(r"!\[(.+?)\]\((.+?)\)", line)
            if match:
                image_path = ROOT / match.group(2)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(5)
                p.add_run().add_picture(str(image_path), width=Inches(6.15 if chinese else 6.75))
                cap = doc.add_paragraph(clean_inline(match.group(1)))
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(8)
                for run in cap.runs:
                    run.font.size = Pt(8.5)
                    run.font.italic = True
                    run.font.color.rgb = RGBColor(*GRAY)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", x) for x in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows, chinese)
            continue
        elif line.startswith("- "):
            doc.add_paragraph(clean_inline(line[2:]), style="List Bullet")
        else:
            p = doc.add_paragraph(clean_inline(line))
            if line.endswith("："):
                p.paragraph_format.keep_with_next = True
            if len(doc.paragraphs) == 2:
                p.paragraph_format.space_after = Pt(17)
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(*GRAY)
        i += 1
    out = ROOT / docx_name
    doc.save(out)
    print(out)


if __name__ == "__main__":
    architecture()
    control_flow()
    legacy_flow()
    total_tokens_chart()
    build("LTGD_Technical_Report_2026-09-30.md", "LTGD_Technical_Report_2026-09-30.docx", False)
    build("LTGD_AI_Use_Report_2026-09-30.md", "LTGD_AI_Use_Report_2026-09-30.docx", True)
